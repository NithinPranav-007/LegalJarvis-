"""Documents and Case Viewer API router with private user uploads and isolation (Task T-14, T-17 / FR-05, FR-08)."""

import datetime
import hashlib
import io
import logging
from pathlib import Path

import docx  # type: ignore[import-untyped]
import pdfplumber  # type: ignore[import-untyped]
from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile, status
from sqlalchemy import delete, or_, select
from sqlalchemy.orm import Session, selectinload

from app.api.auth import get_current_user, get_optional_current_user
from app.core.config import settings
from app.db.models import Chunk, Document, User, UserDocument
from app.db.session import get_sync_db
from app.ingest.embed import get_embedder
from app.ingest.segment import segment_document
from app.retrieval.lexical import ATTRIBUTION_TEXT
from app.schemas.document import ChunkOut, DocumentDetail, UploadResponse, UserDocumentOut

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/v1/documents", tags=["Documents"])

MAX_UPLOAD_BYTES = 20 * 1024 * 1024  # 20 MB per PRD FR-08
ALLOWED_EXTENSIONS = {".pdf", ".txt", ".docx"}


@router.get("/{document_id}", response_model=DocumentDetail)
def get_document(
    document_id: int,
    db: Session = Depends(get_sync_db),
    current_user: User | None = Depends(get_optional_current_user),
) -> DocumentDetail:
    """Retrieve full judgment text, metadata, and ordered chunks for the Case Viewer.

    Enforces private document isolation: If the document is owned by a specific user,
    non-owners receive a 404 Not Found response (preventing cross-user document enumeration).
    """
    stmt = select(Document).options(selectinload(Document.court)).where(Document.id == document_id)
    doc = db.execute(stmt).scalar_one_or_none()
    if not doc:
        raise HTTPException(status_code=404, detail=f"Document with ID {document_id} not found")

    # Isolation check: private documents can only be accessed by their owner (or an admin)
    if doc.owner_user_id is not None:
        if current_user is None or (doc.owner_user_id != current_user.id and current_user.role != "admin"):
            raise HTTPException(status_code=404, detail=f"Document with ID {document_id} not found")

    chunk_stmt = (
        select(Chunk)
        .where(Chunk.document_id == document_id)
        .order_by(Chunk.para_no.asc().nulls_last(), Chunk.start_char.asc())
    )
    chunks = db.execute(chunk_stmt).scalars().all()

    chunks_out = [
        ChunkOut(
            id=c.id,
            para_no=c.para_no,
            segment_type=c.segment_type,
            text=c.text,
            start_char=c.start_char,
            end_char=c.end_char,
        )
        for c in chunks
    ]

    # Resolve document text: from text_path file if it exists, or join chunks
    full_text = ""
    if doc.text_path and Path(doc.text_path).is_file():
        try:
            full_text = Path(doc.text_path).read_text(encoding="utf-8")
        except Exception:
            full_text = "\n\n".join(c.text for c in chunks)
    elif chunks:
        full_text = "\n\n".join(c.text for c in chunks)

    raw_meta = doc.raw_metadata or {}
    bench = raw_meta.get("bench")
    dataset_name = raw_meta.get("dataset_name", "user_upload" if doc.owner_user_id else "sc_2023")

    return DocumentDetail(
        id=doc.id,
        title=doc.title or f"Judgment #{doc.id}",
        court_code=doc.court.code if doc.court else "private",
        court_name=doc.court.name if doc.court else "Private User Document",
        decision_date=doc.decision_date.isoformat() if doc.decision_date else None,
        neutral_citation=doc.neutral_citation,
        bench=bench,
        dataset_name=dataset_name,
        extraction_quality=doc.extraction_quality,
        text=full_text,
        chunks=chunks_out,
        attribution=ATTRIBUTION_TEXT if not doc.owner_user_id else "Uploaded User Document (Private & Confidential)",
    )


@router.get("", response_model=list[dict])
def list_documents(
    limit: int = Query(default=20, ge=1, le=100),
    db: Session = Depends(get_sync_db),
    current_user: User | None = Depends(get_optional_current_user),
) -> list[dict]:
    """List recent judgments for browsing, filtering by user namespace if authenticated."""
    stmt = select(Document).options(selectinload(Document.court))
    if current_user is None:
        stmt = stmt.filter(Document.owner_user_id.is_(None))
    else:
        stmt = stmt.filter(
            or_(Document.owner_user_id.is_(None), Document.owner_user_id == current_user.id)
        )

    stmt = stmt.order_by(Document.id.desc()).limit(limit)
    docs = db.execute(stmt).scalars().all()
    return [
        {
            "id": d.id,
            "title": d.title or f"Judgment #{d.id}",
            "court_code": d.court.code if d.court else "private",
            "court_name": d.court.name if d.court else "Private Document",
            "decision_date": d.decision_date.isoformat() if d.decision_date else None,
            "neutral_citation": d.neutral_citation,
            "is_private": d.owner_user_id is not None,
        }
        for d in docs
    ]


@router.post("/upload", response_model=UploadResponse, status_code=status.HTTP_201_CREATED)
async def upload_user_document(
    file: UploadFile = File(...),
    db: Session = Depends(get_sync_db),
    current_user: User = Depends(get_current_user),
) -> UploadResponse:
    """Upload a private user document (PDF, TXT, DOCX) into an isolated namespace.

    Per PRD FR-08:
    - Maximum size: 20 MB (HTTP 413 on exceed).
    - Formats: PDF, TXT, DOCX (HTTP 415 on unsupported).
    - Unreadable or empty content returns HTTP 422.
    - Chunks are stored with owner_user_id = current_user.id.
    - Retention period: 30 days.
    - LOGS CONTAIN NO DOCUMENT TEXT.
    """
    filename = file.filename or "uploaded_document"
    ext = Path(filename).suffix.lower()

    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE,
            detail=f"Unsupported media type '{ext}'. Allowed formats: {sorted(ALLOWED_EXTENSIONS)}",
        )

    file_bytes = await file.read()
    if len(file_bytes) > MAX_UPLOAD_BYTES:
        raise HTTPException(
            status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            detail=f"File exceeds maximum allowed upload size of {MAX_UPLOAD_BYTES // (1024*1024)} MB",
        )

    sha256 = hashlib.sha256(file_bytes).hexdigest()

    # Text extraction
    extracted_text = ""
    try:
        if ext == ".txt":
            extracted_text = file_bytes.decode("utf-8", errors="replace")
        elif ext == ".pdf":
            with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
                pages_text = []
                for p in pdf.pages:
                    pt = p.extract_text()
                    if pt:
                        pages_text.append(pt)
                extracted_text = "\n\n".join(pages_text)
        elif ext == ".docx":
            doc_file = docx.Document(io.BytesIO(file_bytes))
            extracted_text = "\n\n".join(p.text for p in doc_file.paragraphs if p.text.strip())
    except Exception as e:
        logger.warning("Failed extracting text from uploaded file: %s", type(e).__name__)
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="Failed to read file contents; file may be corrupted or unreadable.",
        ) from e

    cleaned_text = extracted_text.strip()
    if not cleaned_text:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="Scanned or unreadable document: text extraction yielded no extractable content.",
        )

    # Persist file securely in user namespace
    upload_dir = Path("data/uploads") / str(current_user.id)
    upload_dir.mkdir(parents=True, exist_ok=True)
    safe_name = f"{sha256[:16]}_{Path(filename).name}"
    storage_path = upload_dir / safe_name
    storage_path.write_bytes(file_bytes)

    # Segment into paragraphs
    chunk_records = segment_document(cleaned_text)

    # Create Document record
    expires_at = datetime.datetime.now(datetime.UTC) + datetime.timedelta(days=30)
    doc = Document(
        run_id=None,
        court_id=None,
        external_id=f"user_{current_user.id}_{sha256[:12]}",
        title=filename,
        language="en",
        source_file_path=str(storage_path),
        source_sha256=sha256,
        text_path=None,
        extraction_status="ok",
        extraction_quality=1.0,
        raw_metadata={"dataset_name": "user_uploads", "user_id": current_user.id},
        owner_user_id=current_user.id,
    )
    db.add(doc)
    db.flush()

    # Create UserDocument record for retention tracking
    user_doc = UserDocument(
        user_id=current_user.id,
        filename=filename,
        sha256=sha256,
        expires_at=expires_at,
        status="processed",
        document_id=doc.id,
    )
    db.add(user_doc)

    # Generate embeddings and persist chunks
    embedder = get_embedder()
    for cr in chunk_records:
        emb = embedder.embed_text(cr.text) if embedder else None
        chunk = Chunk(
            document_id=doc.id,
            para_no=cr.para_no,
            segment_type=cr.segment_type,
            start_char=cr.start_char,
            end_char=cr.end_char,
            text=cr.text,
            embedding=emb,
            embed_model=settings.EMBED_MODEL if emb else None,
            owner_user_id=current_user.id,
            synthetic=False,
        )
        db.add(chunk)

    db.commit()

    # NOTE: Safe logging only — NO document text in logs
    logger.info(
        "User document uploaded successfully: user_id=%d, doc_id=%d, sha256=%s, chunks=%d, size_bytes=%d",
        current_user.id,
        doc.id,
        sha256,
        len(chunk_records),
        len(file_bytes),
    )

    return UploadResponse(
        document_id=doc.id,
        filename=filename,
        sha256=sha256,
        chunk_count=len(chunk_records),
        expires_at=expires_at.isoformat(),
        message="Document successfully uploaded and processed into private namespace.",
    )


@router.delete("/{document_id}", status_code=status.HTTP_200_OK)
def delete_document(
    document_id: int,
    db: Session = Depends(get_sync_db),
    current_user: User = Depends(get_current_user),
) -> dict:
    """Delete a user document, removing file on disk, chunks, and embeddings."""
    stmt = select(Document).where(Document.id == document_id)
    doc = db.execute(stmt).scalar_one_or_none()

    if not doc:
        raise HTTPException(status_code=404, detail=f"Document with ID {document_id} not found")

    # Protection: Cannot delete public corpus judgments unless admin
    if doc.owner_user_id is None and current_user.role != "admin":
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Cannot delete public corpus judgments.",
        )

    # Isolation check: cannot delete another user's document
    if doc.owner_user_id != current_user.id and current_user.role != "admin":
        raise HTTPException(status_code=404, detail=f"Document with ID {document_id} not found")

    # 1. Remove physical files
    if doc.source_file_path and Path(doc.source_file_path).is_file():
        try:
            Path(doc.source_file_path).unlink(missing_ok=True)
        except Exception as e:
            logger.warning("Error deleting source file %s: %s", doc.source_file_path, e)

    if doc.text_path and Path(doc.text_path).is_file():
        try:
            Path(doc.text_path).unlink(missing_ok=True)
        except Exception as e:
            logger.warning("Error deleting text file %s: %s", doc.text_path, e)

    # 2. Delete UserDocument records
    db.execute(
        delete(UserDocument).where(
            UserDocument.user_id == current_user.id,
            UserDocument.sha256 == doc.source_sha256,
        )
    )

    # 3. Delete Chunks and Document
    db.execute(delete(Chunk).where(Chunk.document_id == document_id))
    db.delete(doc)
    db.commit()

    logger.info("User document deleted: user_id=%d, doc_id=%d", current_user.id, document_id)
    return {"status": "deleted", "document_id": document_id}


@router.get("/user/list", response_model=list[UserDocumentOut])
def list_user_documents(
    db: Session = Depends(get_sync_db),
    current_user: User = Depends(get_current_user),
) -> list[UserDocumentOut]:
    """List all private documents uploaded by the authenticated user."""
    stmt = (
        select(UserDocument)
        .where(UserDocument.user_id == current_user.id)
        .order_by(UserDocument.id.desc())
    )
    udocs = db.execute(stmt).scalars().all()
    return [
        UserDocumentOut(
            id=u.id,
            document_id=u.document_id,
            filename=u.filename,
            sha256=u.sha256,
            status=u.status,
            expires_at=u.expires_at.isoformat(),
        )
        for u in udocs
    ]
